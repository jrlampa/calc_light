import io
import msoffcrypto
import openpyxl
import re
from typing import Dict, Any, List

def normalize_name(name: Any) -> str:
    """Remove espaços extras, caracteres invisíveis e normaliza para string."""
    if name is None: return ""
    # Remove caracteres Unicode invisíveis (ex: \xa0)
    s = str(name).replace('\xa0', ' ').strip()
    # Remove espaços duplos
    s = re.sub(r'\s+', ' ', s)
    return s

class LegacyImporter:
    """
    Classe para extração de dados de planilhas Excel (.xlsm) protegidas.
    Utiliza msoffcrypto-tool e openpyxl para garantir compatibilidade cross-platform (Linux/Docker).
    """
    
    def __init__(self, file_path: str, password: str = None):
        self.file_path = file_path
        self.password = password
        self.workbook = None

    def open_workbook(self):
        """Abre a planilha, descriptografando se necessário."""
        try:
            office_file = open(self.file_path, "rb")
            file = msoffcrypto.OfficeFile(office_file)
            
            if file.is_encrypted():
                if self.password:
                    file.load_key(password=self.password)
                else:
                    # Tenta senha padrão se estiver protegido mas sem senha informada
                    # Algumas planilhas usam 'VelvetSweatshop' como senha padrão
                    try:
                        file.load_key(password='VelvetSweatshop')
                    except:
                        raise ValueError("Arquivo criptografado e nenhuma senha válida fornecida.")
                
                decrypted = io.BytesIO()
                file.decrypt(decrypted)
                decrypted.seek(0)
                self.workbook = openpyxl.load_workbook(decrypted, data_only=True)
            else:
                # Não está criptografado, abre direto
                self.workbook = openpyxl.load_workbook(self.file_path, data_only=True)
        except Exception as e:
            # Fallback final direto caso msoffcrypto falhe na detecção
            try:
                self.workbook = openpyxl.load_workbook(self.file_path, data_only=True)
            except:
                raise e
        return self.workbook

    def extract_node_data(self, sheet_name: str) -> Dict[str, Any]:
        """Extrai os dados de entrada e resultados de uma aba 'Ponto (X)'."""
        if not self.workbook:
            self.open_workbook()
            
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"Aba {sheet_name} não encontrada.")
            
        ws = self.workbook[sheet_name]
        
        data = {
            "metadata": {
                "project": ws["H1"].value,
                "point_id": ws["K1"].value,
                "pole_type": ws["C7"].value,
                "pole_model": ws["C8"].value
            },
            "inputs": {
                "mt1": {
                    "t1": {"network": normalize_name(ws["C12"].value), "cable": normalize_name(ws["C13"].value), "span": ws["C14"].value, "sag": ws["C15"].value, "angle": ws["C16"].value, "cable_qty": ws["C21"].value},
                    "t2": {"network": normalize_name(ws["F12"].value), "cable": normalize_name(ws["F13"].value), "span": ws["F14"].value, "sag": ws["F15"].value, "angle": ws["F16"].value, "cable_qty": ws["F21"].value},
                    "t3": {"network": normalize_name(ws["I12"].value), "cable": normalize_name(ws["I13"].value), "span": ws["I14"].value, "sag": ws["I15"].value, "angle": ws["I16"].value, "cable_qty": ws["I21"].value},
                    "t4": {"network": normalize_name(ws["L12"].value), "cable": normalize_name(ws["L13"].value), "span": ws["L14"].value, "sag": ws["L15"].value, "angle": ws["L16"].value, "cable_qty": ws["L21"].value},
                },
                "pole_height": ws["C17"].value,
                "anchorage_height": ws["C18"].value
            },
            "expected_results": {
                "resultante_mt1_daN": ws["C32"].value,
                "tracao_mt1_topo_daN": ws["F33"].value,
                "resultante_total_daN": ws["C140"].value,
                "angulo_total_deg": ws["D141"].value,
                "vento_poste_daN": ws["C149"].value
            }
        }
        return data

    def list_calculation_sheets(self) -> List[str]:
        """Lista todas as abas que contêm cálculos de pontos."""
        if not self.workbook:
            self.open_workbook()
        return [s for s in self.workbook.sheetnames if s.startswith("Ponto (")]

if __name__ == "__main__":
    # Teste rápido de extração
    import os
    import pathlib
    # Repo root is 3 levels up: scripts/ -> app/ -> backend/ -> repo root
    xlsm_path = str(pathlib.Path(__file__).parents[3] / "CÁLCULO DE TRAÇÃO OII-25-2249.xlsm")
    if os.path.exists(xlsm_path):
        importer = LegacyImporter(xlsm_path)
        importer.open_workbook()
        points = importer.list_calculation_sheets()
        print(f"Abas encontradas: {points}")
        if points:
            sample = importer.extract_node_data(points[0])
            print(f"Dados extraídos de {points[0]}: {sample['expected_results']}")
