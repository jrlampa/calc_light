"""Testes paranóicos para o motor de parsing GIS.

Cobertura exigida: 100% do módulo app.domain.gis_parser.

Cenários cobertos por formato:
- KML: válido, XML malformado, sem placemarks, coordenadas ausentes, sem namespace,
       nome ausente (usa default "Ponto"), coordenadas com formato inválido
- KMZ: válido, ZIP corrompido, ZIP sem .kml interno
- GeoJSON: válido, JSON inválido, JSON não-dict, features não-lista,
           geometria não-Point, coordenadas ausentes/inválidas, sem prop name
- Excel: válido, bytes inválidos, planilha vazia, colunas lat/lng ausentes,
         linha com valores não-numéricos (deve ser ignorada), aba ativa None (mock)
- parse_file: dispatch correto para cada extensão, extensão não suportada,
              arquivo sem extensão
"""

from __future__ import annotations

import io
import xml.etree.ElementTree as ET
import zipfile
from unittest.mock import MagicMock, patch

import openpyxl
import pytest

from app.domain.gis_parser import (
    _find_col,
    _kml_coords_to_latlon,
    _parse_kml_tree,
    parse_excel,
    parse_file,
    parse_geojson,
    parse_kml,
    parse_kmz,
)

# ─── Helpers de fixture ───────────────────────────────────────────────────────


def _make_kml(placemarks: list[dict]) -> bytes:
    """Gera bytes de um KML 2.2 com os placemarks informados."""
    inner = ""
    for pm in placemarks:
        name_part = f"<name>{pm.get('name', '')}</name>" if "name" in pm else ""
        coord_part = (
            f"<Point><coordinates>{pm['coords']}</coordinates></Point>"
            if "coords" in pm
            else ""
        )
        inner += f"<Placemark>{name_part}{coord_part}</Placemark>"
    return (
        f'<?xml version="1.0"?>'
        f'<kml xmlns="http://www.opengis.net/kml/2.2"><Document>{inner}</Document></kml>'
    ).encode()


def _make_kml_no_ns(placemarks: list[dict]) -> bytes:
    """Gera bytes de KML sem namespace."""
    inner = ""
    for pm in placemarks:
        name_part = f"<name>{pm.get('name', '')}</name>" if "name" in pm else ""
        coord_part = (
            f"<Point><coordinates>{pm['coords']}</coordinates></Point>"
            if "coords" in pm
            else ""
        )
        inner += f"<Placemark>{name_part}{coord_part}</Placemark>"
    return f"<kml><Document>{inner}</Document></kml>".encode()


def _make_kmz(kml_bytes: bytes, kml_name: str = "doc.kml") -> bytes:
    """Empacota bytes KML dentro de um ZIP (.kmz)."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(kml_name, kml_bytes)
    return buf.getvalue()


def _make_geojson(features: list[dict]) -> bytes:
    """Gera bytes GeoJSON com a lista de features fornecida."""
    import json
    return json.dumps({"type": "FeatureCollection", "features": features}).encode()


def _make_excel_bytes(rows: list[list]) -> bytes:
    """Gera bytes de um .xlsx com as linhas fornecidas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


# ─── _kml_coords_to_latlon ────────────────────────────────────────────────────


def test_kml_coords_to_latlon_valid():
    lat, lng = _kml_coords_to_latlon("-43.5,-22.9,0")
    assert lat == pytest.approx(-22.9)
    assert lng == pytest.approx(-43.5)


def test_kml_coords_to_latlon_two_parts():
    lat, lng = _kml_coords_to_latlon("-43.5,-22.9")
    assert lat == pytest.approx(-22.9)
    assert lng == pytest.approx(-43.5)


def test_kml_coords_to_latlon_single_part_raises():
    with pytest.raises(ValueError, match="inválida"):
        _kml_coords_to_latlon("-43.5")


# ─── _find_col ────────────────────────────────────────────────────────────────


def test_find_col_found():
    assert _find_col(["name", "lat", "lng"], ["lat", "latitude"]) == 1


def test_find_col_not_found():
    assert _find_col(["name", "x", "y"], ["lat", "latitude"]) is None


# ─── _parse_kml_tree ─────────────────────────────────────────────────────────


def test_parse_kml_tree_with_namespace():
    kml = _make_kml([{"name": "P1", "coords": "-43.0,-22.0,0"}])
    root = ET.fromstring(kml.decode())
    points = _parse_kml_tree(root)
    assert len(points) == 1
    assert points[0].label == "P1"
    assert points[0].lat == pytest.approx(-22.0)


def test_parse_kml_tree_without_namespace():
    kml = _make_kml_no_ns([{"name": "P2", "coords": "-44.0,-23.0"}])
    root = ET.fromstring(kml.decode())
    points = _parse_kml_tree(root)
    assert len(points) == 1
    assert points[0].label == "P2"


def test_parse_kml_tree_no_placemarks():
    root = ET.fromstring("<kml><Document></Document></kml>")
    assert _parse_kml_tree(root) == []


def test_parse_kml_tree_missing_name_defaults_to_ponto():
    kml = _make_kml([{"coords": "-43.0,-22.0,0"}])  # no "name" key
    root = ET.fromstring(kml.decode())
    points = _parse_kml_tree(root)
    assert points[0].label == "Ponto"


def test_parse_kml_tree_missing_coordinates_skipped():
    kml = _make_kml([{"name": "No Coords"}])  # no "coords" key
    root = ET.fromstring(kml.decode())
    assert _parse_kml_tree(root) == []


def test_parse_kml_tree_invalid_coords_skipped():
    kml = _make_kml([{"name": "Bad", "coords": "only_one_value"}])
    root = ET.fromstring(kml.decode())
    assert _parse_kml_tree(root) == []


# ─── parse_kml ───────────────────────────────────────────────────────────────


def test_parse_kml_valid():
    kml = _make_kml([
        {"name": "Alpha", "coords": "-43.1,-22.1,0"},
        {"name": "Beta", "coords": "-43.2,-22.2,0"},
    ])
    points = parse_kml(kml)
    assert len(points) == 2
    assert points[0].label == "Alpha"
    assert points[1].label == "Beta"


def test_parse_kml_malformed_xml_raises():
    with pytest.raises(ValueError, match="KML inválido"):
        parse_kml(b"<not valid xml")


def test_parse_kml_empty_document_returns_empty_list():
    kml = b'<kml xmlns="http://www.opengis.net/kml/2.2"><Document></Document></kml>'
    assert parse_kml(kml) == []


# ─── parse_kmz ───────────────────────────────────────────────────────────────


def test_parse_kmz_valid():
    kml = _make_kml([{"name": "KMZ Point", "coords": "-43.0,-22.0,0"}])
    kmz = _make_kmz(kml)
    points = parse_kmz(kmz)
    assert len(points) == 1
    assert points[0].label == "KMZ Point"


def test_parse_kmz_bad_zip_raises():
    with pytest.raises(ValueError, match="não é um arquivo ZIP"):
        parse_kmz(b"this is not a zip file")


def test_parse_kmz_no_kml_inside_raises():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.txt", "hello")
    with pytest.raises(ValueError, match="Nenhum arquivo .kml"):
        parse_kmz(buf.getvalue())


def test_parse_kmz_nested_kml_name():
    """KMZ pode conter KML com nome de subpasta, ex: 'doc/doc.kml'."""
    kml = _make_kml([{"name": "Sub", "coords": "-43.0,-22.0,0"}])
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("folder/map.kml", kml)
    points = parse_kmz(buf.getvalue())
    assert len(points) == 1


# ─── parse_geojson ───────────────────────────────────────────────────────────


def test_parse_geojson_valid_with_name():
    gj = _make_geojson([
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [-43.0, -22.0]},
         "properties": {"name": "GJ-P1"}}
    ])
    points = parse_geojson(gj)
    assert len(points) == 1
    assert points[0].label == "GJ-P1"
    assert points[0].lat == pytest.approx(-22.0)
    assert points[0].lng == pytest.approx(-43.0)


def test_parse_geojson_valid_with_label_property():
    gj = _make_geojson([
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [-43.0, -22.0]},
         "properties": {"label": "LBL"}}
    ])
    points = parse_geojson(gj)
    assert points[0].label == "LBL"


def test_parse_geojson_valid_no_name_defaults_to_ponto():
    gj = _make_geojson([
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [-43.0, -22.0]},
         "properties": {}}
    ])
    points = parse_geojson(gj)
    assert points[0].label == "Ponto"


def test_parse_geojson_invalid_json_raises():
    with pytest.raises(ValueError, match="GeoJSON inválido"):
        parse_geojson(b"{invalid json}")


def test_parse_geojson_not_dict_raises():
    import json
    with pytest.raises(ValueError, match="objeto JSON"):
        parse_geojson(json.dumps([1, 2, 3]).encode())


def test_parse_geojson_features_not_list_raises():
    import json
    with pytest.raises(ValueError, match="'features'"):
        parse_geojson(json.dumps({"features": "bad"}).encode())


def test_parse_geojson_no_features_key_returns_empty():
    import json
    points = parse_geojson(json.dumps({"type": "FeatureCollection"}).encode())
    assert points == []


def test_parse_geojson_non_point_geometry_skipped():
    gj = _make_geojson([
        {"type": "Feature", "geometry": {"type": "LineString", "coordinates": [[-43.0, -22.0]]},
         "properties": {"name": "Line"}}
    ])
    assert parse_geojson(gj) == []


def test_parse_geojson_missing_coordinates_skipped():
    import json
    gj = json.dumps({"type": "FeatureCollection", "features": [
        {"type": "Feature", "geometry": {"type": "Point"}, "properties": {"name": "X"}}
    ]}).encode()
    assert parse_geojson(gj) == []


def test_parse_geojson_coordinates_too_short_skipped():
    gj = _make_geojson([
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [-43.0]},
         "properties": {"name": "X"}}
    ])
    assert parse_geojson(gj) == []


def test_parse_geojson_bad_coordinate_values_skipped():
    gj = _make_geojson([
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": ["bad", "value"]},
         "properties": {"name": "X"}}
    ])
    assert parse_geojson(gj) == []


def test_parse_geojson_non_dict_feature_skipped():
    import json
    gj = json.dumps({"type": "FeatureCollection", "features": ["not_a_dict"]}).encode()
    assert parse_geojson(gj) == []


def test_parse_geojson_null_geometry_skipped():
    import json
    gj = json.dumps({"type": "FeatureCollection", "features": [
        {"type": "Feature", "geometry": None, "properties": {"name": "X"}}
    ]}).encode()
    assert parse_geojson(gj) == []


# ─── parse_excel ─────────────────────────────────────────────────────────────


def test_parse_excel_valid_all_columns():
    content = _make_excel_bytes([
        ["Label", "Lat", "Lng"],
        ["P1", -22.9, -43.2],
        ["P2", -23.0, -43.3],
    ])
    points = parse_excel(content)
    assert len(points) == 2
    assert points[0].label == "P1"
    assert points[0].lat == pytest.approx(-22.9)
    assert points[0].lng == pytest.approx(-43.2)


def test_parse_excel_with_latitude_longitude_columns():
    content = _make_excel_bytes([
        ["Name", "Latitude", "Longitude"],
        ["X1", -22.5, -43.5],
    ])
    points = parse_excel(content)
    assert len(points) == 1
    assert points[0].label == "X1"


def test_parse_excel_no_label_column_defaults_to_ponto():
    content = _make_excel_bytes([
        ["Lat", "Lng"],
        [-22.9, -43.2],
    ])
    points = parse_excel(content)
    assert points[0].label == "Ponto"


def test_parse_excel_empty_label_cell_defaults_to_ponto():
    content = _make_excel_bytes([
        ["Label", "Lat", "Lng"],
        ["", -22.9, -43.2],
    ])
    points = parse_excel(content)
    assert points[0].label == "Ponto"


def test_parse_excel_bad_rows_skipped():
    """Linhas com valores não numéricos em lat/lng devem ser silenciosamente ignoradas."""
    content = _make_excel_bytes([
        ["Label", "Lat", "Lng"],
        ["P1", "not_a_number", -43.2],
        ["P2", -22.9, -43.3],
    ])
    points = parse_excel(content)
    assert len(points) == 1
    assert points[0].label == "P2"


def test_parse_excel_missing_lat_lng_cols_raises():
    content = _make_excel_bytes([
        ["Name", "X", "Y"],
        ["A", 1, 2],
    ])
    with pytest.raises(ValueError, match="lat"):
        parse_excel(content)


def test_parse_excel_only_header_returns_empty():
    content = _make_excel_bytes([["Label", "Lat", "Lng"]])
    points = parse_excel(content)
    assert points == []


def test_parse_excel_invalid_bytes_raises():
    with pytest.raises(ValueError, match="Excel inválido"):
        parse_excel(b"this is not an excel file at all")


def test_parse_excel_no_active_sheet_raises():
    """Quando wb.active é None, deve levantar ValueError com 'aba ativa'."""
    with patch("app.domain.gis_parser.openpyxl.load_workbook") as mock_load:
        mock_wb = MagicMock()
        mock_wb.active = None
        mock_wb.close = MagicMock()
        mock_load.return_value = mock_wb
        with pytest.raises(ValueError, match="aba ativa"):
            parse_excel(b"fake")


def test_parse_excel_empty_sheet_raises():
    """Planilha sem nenhuma linha (nem cabeçalho) deve levantar ValueError."""
    with patch("app.domain.gis_parser.openpyxl.load_workbook") as mock_load:
        mock_ws = MagicMock()
        mock_ws.iter_rows.return_value = iter([])
        mock_wb = MagicMock()
        mock_wb.active = mock_ws
        mock_wb.close = MagicMock()
        mock_load.return_value = mock_wb
        with pytest.raises(ValueError, match="vazia"):
            parse_excel(b"fake")


# ─── parse_file (dispatcher) ─────────────────────────────────────────────────


def test_parse_file_dispatch_kml():
    kml = _make_kml([{"name": "D", "coords": "-43.0,-22.0,0"}])
    points = parse_file("mapa.kml", kml)
    assert len(points) == 1


def test_parse_file_dispatch_kmz():
    kml = _make_kml([{"name": "D", "coords": "-43.0,-22.0,0"}])
    kmz = _make_kmz(kml)
    points = parse_file("mapa.kmz", kmz)
    assert len(points) == 1


def test_parse_file_dispatch_geojson():
    gj = _make_geojson([
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [-43.0, -22.0]},
         "properties": {"name": "D"}}
    ])
    points = parse_file("mapa.geojson", gj)
    assert len(points) == 1


def test_parse_file_dispatch_json():
    gj = _make_geojson([
        {"type": "Feature", "geometry": {"type": "Point", "coordinates": [-43.0, -22.0]},
         "properties": {"name": "D"}}
    ])
    points = parse_file("mapa.json", gj)
    assert len(points) == 1


def test_parse_file_dispatch_xlsx():
    content = _make_excel_bytes([["Label", "Lat", "Lng"], ["P", -22.9, -43.2]])
    points = parse_file("mapa.xlsx", content)
    assert len(points) == 1


def test_parse_file_dispatch_xls_alias():
    """Extensão .xls deve ser tratada como .xlsx (alias)."""
    content = _make_excel_bytes([["Label", "Lat", "Lng"], ["P", -22.9, -43.2]])
    points = parse_file("mapa.xls", content)
    assert len(points) == 1


def test_parse_file_unsupported_extension_raises():
    with pytest.raises(ValueError, match="não suportada"):
        parse_file("mapa.csv", b"data")


def test_parse_file_no_extension_raises():
    with pytest.raises(ValueError, match="não suportada"):
        parse_file("mapafile", b"data")


def test_parse_file_case_insensitive_extension():
    kml = _make_kml([{"name": "D", "coords": "-43.0,-22.0,0"}])
    points = parse_file("mapa.KML", kml)
    assert len(points) == 1
