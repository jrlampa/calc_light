"""Testes paranóicos para o motor de exportação Excel em lotes.

Cenário principal: projeto com 35 postes
- Lote_01/ contém poste_01.xlsm … poste_30.xlsm
- Lote_02/ contém poste_31.xlsm … poste_35.xlsm
- Célula C8 ("Modelo do Poste") do poste_35 deve conter o valor injetado.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from app.domain.excel_exporter import BATCH_SIZE, build_export_zip, build_poste_xlsm
from app.domain.excel_mapping import CALC_SHEET, CELL_MAP

TEMPLATE = Path(__file__).parent.parent / "templates" / "modelo.xlsm"


# ─── Fixtures ────────────────────────────────────────────────────────────────


def _make_poste_data(idx: int) -> dict[str, Any]:
    """Gera um dicionário de dados mínimo para o poste de índice ``idx``."""
    return {
        "projeto": f"Projeto Teste QA",
        "ponto": idx,
        "tipo_poste": "Concreto circular",
        "modelo_poste": f"11 m / {idx * 10} daN",
        "mt1_t1_rede": "Compacta",
        "mt1_t1_cabo": "1/0AWG-CAA, XLPE, 13,8 kV",
        "mt1_t1_vao": float(50 + idx),
        "mt1_t1_flecha": 0.5,
        "mt1_t1_angulo": 0.0,
        "mt1_altura_poste": 9.3,
        "mt1_altura_ancoragem": 9.1,
    }


@pytest.fixture(scope="module")
def postes_35() -> list[dict[str, Any]]:
    return [_make_poste_data(i) for i in range(1, 36)]


@pytest.fixture(scope="module")
def zip_35(postes_35: list[dict[str, Any]]) -> bytes:
    return build_export_zip(postes_35, template_path=TEMPLATE)


# ─── Testes de sanidade do motor ─────────────────────────────────────────────


def test_batch_size_constant():
    """BATCH_SIZE deve ser exatamente 30."""
    assert BATCH_SIZE == 30


def test_cell_map_has_required_keys():
    """CELL_MAP deve conter pelo menos as células críticas de MT1."""
    required = {
        "modelo_poste", "tipo_poste", "projeto", "ponto",
        "mt1_t1_rede", "mt1_t1_cabo", "mt1_t1_vao",
        "mt1_t1_flecha", "mt1_t1_angulo",
        "mt1_altura_poste", "mt1_altura_ancoragem",
    }
    assert required.issubset(CELL_MAP.keys())


def test_cell_map_modelo_poste_address():
    """A célula de 'modelo_poste' deve ser C8."""
    assert CELL_MAP["modelo_poste"] == "C8"


# ─── Testes de injeção de dados (single file) ─────────────────────────────────


def test_build_poste_xlsm_returns_bytes():
    xlsm = build_poste_xlsm(_make_poste_data(1), template_path=TEMPLATE)
    assert isinstance(xlsm, bytes)
    assert len(xlsm) > 0


def test_build_poste_xlsm_injects_modelo_poste():
    data = _make_poste_data(7)
    xlsm = build_poste_xlsm(data, template_path=TEMPLATE)

    wb = openpyxl.load_workbook(io.BytesIO(xlsm), keep_vba=True)
    ws = wb[CALC_SHEET]
    assert ws["C8"].value == data["modelo_poste"]


def test_build_poste_xlsm_injects_ponto():
    data = _make_poste_data(42)
    xlsm = build_poste_xlsm(data, template_path=TEMPLATE)

    wb = openpyxl.load_workbook(io.BytesIO(xlsm), keep_vba=True)
    ws = wb[CALC_SHEET]
    assert ws[CELL_MAP["ponto"]].value == 42


def test_build_poste_xlsm_injects_mt1_vao():
    data = _make_poste_data(3)
    xlsm = build_poste_xlsm(data, template_path=TEMPLATE)

    wb = openpyxl.load_workbook(io.BytesIO(xlsm), keep_vba=True)
    ws = wb[CALC_SHEET]
    assert ws[CELL_MAP["mt1_t1_vao"]].value == pytest.approx(data["mt1_t1_vao"])


def test_build_poste_xlsm_ignores_missing_keys():
    """Chaves ausentes não devem lançar exceção."""
    data = {"modelo_poste": "Só modelo"}
    xlsm = build_poste_xlsm(data, template_path=TEMPLATE)
    assert isinstance(xlsm, bytes)


# ─── Testes de ZIP com 35 postes ──────────────────────────────────────────────


def test_zip_35_is_valid_zip(zip_35: bytes):
    assert zipfile.is_zipfile(io.BytesIO(zip_35))


def test_zip_35_total_file_count(zip_35: bytes):
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        xlsm_files = [n for n in zf.namelist() if n.endswith(".xlsm")]
    assert len(xlsm_files) == 35


def test_zip_35_has_lote_01_folder(zip_35: bytes):
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        names = zf.namelist()
    assert any(n.startswith("Lote_01/") for n in names)


def test_zip_35_has_lote_02_folder(zip_35: bytes):
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        names = zf.namelist()
    assert any(n.startswith("Lote_02/") for n in names)


def test_zip_35_no_lote_03_folder(zip_35: bytes):
    """Com apenas 35 postes não deve existir Lote_03."""
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        names = zf.namelist()
    assert not any(n.startswith("Lote_03/") for n in names)


def test_zip_35_lote_01_has_30_files(zip_35: bytes):
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        lote1 = [n for n in zf.namelist() if n.startswith("Lote_01/") and n.endswith(".xlsm")]
    assert len(lote1) == 30


def test_zip_35_lote_02_has_5_files(zip_35: bytes):
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        lote2 = [n for n in zf.namelist() if n.startswith("Lote_02/") and n.endswith(".xlsm")]
    assert len(lote2) == 5


def test_zip_35_lote_01_first_file(zip_35: bytes):
    """Primeiro arquivo do Lote_01 deve ser poste_01.xlsm."""
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        lote1 = sorted(n for n in zf.namelist() if n.startswith("Lote_01/") and n.endswith(".xlsm"))
    assert lote1[0] == "Lote_01/poste_01.xlsm"


def test_zip_35_lote_01_last_file(zip_35: bytes):
    """Último arquivo do Lote_01 deve ser poste_30.xlsm."""
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        lote1 = sorted(n for n in zf.namelist() if n.startswith("Lote_01/") and n.endswith(".xlsm"))
    assert lote1[-1] == "Lote_01/poste_30.xlsm"


def test_zip_35_lote_02_last_file(zip_35: bytes):
    """Último arquivo do Lote_02 deve ser poste_35.xlsm."""
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        lote2 = sorted(n for n in zf.namelist() if n.startswith("Lote_02/") and n.endswith(".xlsm"))
    assert lote2[-1] == "Lote_02/poste_35.xlsm"


def test_zip_35_poste_35_cell_c8(zip_35: bytes, postes_35: list[dict[str, Any]]):
    """Célula C8 (modelo_poste) do poste_35 deve conter o valor injetado."""
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        xlsm_bytes = zf.read("Lote_02/poste_35.xlsm")

    wb = openpyxl.load_workbook(io.BytesIO(xlsm_bytes), keep_vba=True)
    ws = wb[CALC_SHEET]
    expected = postes_35[34]["modelo_poste"]  # índice 34 → poste 35
    assert ws["C8"].value == expected


def test_zip_35_poste_35_cell_ponto(zip_35: bytes):
    """Célula K1 (ponto) do poste_35 deve ser 35."""
    with zipfile.ZipFile(io.BytesIO(zip_35)) as zf:
        xlsm_bytes = zf.read("Lote_02/poste_35.xlsm")

    wb = openpyxl.load_workbook(io.BytesIO(xlsm_bytes), keep_vba=True)
    ws = wb[CALC_SHEET]
    assert ws[CELL_MAP["ponto"]].value == 35


# ─── Testes de lote com exatamente 30 postes (sem subpastas) ──────────────────


def test_zip_30_no_lote_folders():
    """Com exatamente 30 postes os arquivos devem ficar na raiz do ZIP."""
    data_30 = [_make_poste_data(i) for i in range(1, 31)]
    zip_bytes = build_export_zip(data_30, template_path=TEMPLATE)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        names = zf.namelist()
    assert not any("Lote_" in n for n in names)
    assert "poste_30.xlsm" in names


def test_zip_31_uses_lote_folders():
    """Com 31 postes deve usar subpastas Lote_01/ e Lote_02/."""
    data_31 = [_make_poste_data(i) for i in range(1, 32)]
    zip_bytes = build_export_zip(data_31, template_path=TEMPLATE)
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        names = zf.namelist()
    assert any(n.startswith("Lote_01/") for n in names)
    assert any(n.startswith("Lote_02/") for n in names)
