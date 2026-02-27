import openpyxl

paths = [
    r"c:\myworld\CACL_LIGHT\CÁLCULO DE TRAÇÃO OII-25-2249.xlsm",
    r"c:\myworld\CACL_LIGHT\POSTE69.xlsm"
]

with open(r"c:\myworld\CACL_LIGHT\excel_dump.txt", "w", encoding="utf-8") as f:
    for p in paths:
        f.write(f"\n====================\n{p}\n====================\n")
        try:
            wb_val = openpyxl.load_workbook(p, data_only=True)
            wb_form = openpyxl.load_workbook(p, data_only=False)
            f.write(f"Sheets: {wb_val.sheetnames}\n")
            
            for sheet in wb_val.sheetnames:
                ws_val = wb_val[sheet]
                ws_form = wb_form[sheet]
                f.write(f"\n--- Sheet '{sheet}' ---\n")
                
                # Zip is not completely reliable if dimensions differ, but openpyxl usually keeps them same.
                for row_v, row_f in zip(ws_val.iter_rows(), ws_form.iter_rows()):
                    for cell_v, cell_f in zip(row_v, row_f):
                        val = cell_v.value
                        form = cell_f.value
                        
                        if val is not None or form is not None:
                            val_str = str(val).strip() if val is not None else ""
                            form_str = str(form).strip() if form is not None else ""
                            
                            if form_str.startswith("=") and val_str != form_str:
                                f.write(f"{cell_v.coordinate}: val=[{val_str}], formula=[{form_str}]\n")
                            elif val_str:
                                f.write(f"{cell_v.coordinate}: {val_str}\n")
        except Exception as e:
            f.write(f"Error reading {p}: {e}\n")
