import openpyxl

paths = [
    r"c:\myworld\CACL_LIGHT\CÁLCULO DE TRAÇÃO OII-25-2249.xlsm",
    r"c:\myworld\CACL_LIGHT\POSTE69.xlsm"
]

for p in paths:
    print(f"\n====================\nAnalyzing {p}\n====================")
    try:
        wb = openpyxl.load_workbook(p, data_only=True)
        print("Sheets:", wb.sheetnames)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            data = []
            for row in ws.iter_rows(min_row=1, max_row=30, values_only=True):
                # Only keep rows that are not entirely None
                if any(x is not None for x in row):
                    # Replace None with empty string for better printing
                    data.append([str(x) if x is not None else "" for x in row])
            if data:
                print(f"\n--- Sheet '{sheet}' has {len(data)} non-empty rows in first 30. Sample:")
                for r in data[:10]:
                    print(" | ".join(r[:10]))  # Print first 10 cols
    except Exception as e:
        print(f"Error reading {p}: {e}")
